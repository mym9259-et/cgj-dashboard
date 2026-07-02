from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.car_series_mapping import CarSeriesMapping
from app.models.mapping_metadata import MappingMetadata
from app.models.personnel_mapping import PersonnelMapping
from app.models.store_mapping import StoreMapping


MAPPING_CONFIG = {
    "store": {
        "model": StoreMapping,
        "key": "merchant_name",
        "columns": ["merchant_name", "lingpao_region", "province", "city", "is_lingpao", "store_manager", "dealer_direct", "store_mode"],
        "labels": ["商户名称", "大区", "省份", "城市", "是否零售", "门店总经理", "经销商/直营", "模式"],
    },
    "car-series": {
        "model": CarSeriesMapping,
        "key": "raw_series",
        "columns": ["raw_series", "clean_series", "brand"],
        "labels": ["原始车系", "车系（修正后）", "品牌"],
    },
    "personnel": {
        "model": PersonnelMapping,
        "key": "salesperson",
        "columns": ["salesperson", "is_active", "role"],
        "labels": ["人员姓名", "是否在职", "角色"],
    },
}


def _config(mapping_type: str) -> dict:
    if mapping_type not in MAPPING_CONFIG:
        raise ValueError("不支持的映射表类型")
    return MAPPING_CONFIG[mapping_type]


def _as_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "是", "在职"}


def _normalize(mapping_type: str, row: dict) -> dict:
    cfg = _config(mapping_type)
    result = {}
    for column in cfg["columns"]:
        value = row.get(column)
        if column == "is_active":
            result[column] = _as_bool(value)
        else:
            result[column] = str(value).strip() if value not in (None, "") else None
    if not result[cfg["key"]]:
        raise ValueError(f"{cfg['key']} 不能为空")
    return result


async def list_rows(db: AsyncSession, mapping_type: str) -> list[dict]:
    cfg = _config(mapping_type)
    model = cfg["model"]
    rows = (await db.scalars(select(model).order_by(getattr(model, cfg["key"])))).all()
    return [{"id": row.id, **{column: getattr(row, column) for column in cfg["columns"]}} for row in rows]


async def replace_rows(db: AsyncSession, mapping_type: str, rows: list[dict]) -> int:
    cfg = _config(mapping_type)
    model = cfg["model"]
    normalized = [_normalize(mapping_type, row) for row in rows]
    keys = [row[cfg["key"]] for row in normalized]
    if len(keys) != len(set(keys)):
        duplicates = sorted({key for key in keys if keys.count(key) > 1})
        raise ValueError(f"关键字段存在重复值：{'、'.join(duplicates[:20])}")
    await db.execute(delete(model))
    db.add_all([model(**row) for row in normalized])
    await db.flush()
    return len(normalized)


async def set_metadata(db: AsyncSession, mapping_type: str, source_type: str, filename: str | None, username: str) -> None:
    metadata = await db.scalar(select(MappingMetadata).where(MappingMetadata.mapping_type == mapping_type))
    if metadata is None:
        metadata = MappingMetadata(mapping_type=mapping_type)
        db.add(metadata)
    metadata.source_type = source_type
    metadata.source_filename = filename
    metadata.updated_by = username
    metadata.updated_at = datetime.now()


async def get_metadata(db: AsyncSession, mapping_type: str) -> dict:
    cfg = _config(mapping_type)
    metadata = await db.scalar(select(MappingMetadata).where(MappingMetadata.mapping_type == mapping_type))
    count = len(await list_rows(db, mapping_type))
    return {
        "mapping_type": mapping_type,
        "columns": [{"key": key, "label": label} for key, label in zip(cfg["columns"], cfg["labels"])],
        "row_count": count,
        "source_type": metadata.source_type if metadata else None,
        "source_filename": metadata.source_filename if metadata else None,
        "updated_by": metadata.updated_by if metadata else None,
        "updated_at": metadata.updated_at if metadata else None,
    }


async def import_personnel_excel(db: AsyncSession, filepath: str) -> int:
    workbook = load_workbook(filepath, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        aliases = {"人员姓名": "salesperson", "姓名": "salesperson", "是否在职": "is_active", "在职状态": "is_active", "角色": "role"}
        rows = []
        for values_row in values:
            raw = {aliases.get(header, header): value for header, value in zip(headers, values_row)}
            if raw.get("salesperson"):
                rows.append(raw)
        return await replace_rows(db, "personnel", rows)
    finally:
        workbook.close()


async def export_workbook(db: AsyncSession, mapping_type: str) -> BytesIO:
    cfg = _config(mapping_type)
    rows = await list_rows(db, mapping_type)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "映射表"
    sheet.append(cfg["labels"])
    for row in rows:
        values = []
        for column in cfg["columns"]:
            value = row.get(column)
            if column == "is_active":
                value = "是" if value else "否"
            values.append(value)
        sheet.append(values)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
